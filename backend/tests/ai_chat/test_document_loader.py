"""Test thuần logic cho document_loader.py — không cần DB, không cần API key.

Dùng chính các thư viện đọc file (python-docx, python-pptx, openpyxl) để tự
sinh file thật trong thư mục tạm (tmp_path của pytest), tránh phụ thuộc vào
fixture file đính kèm sẵn trong repo.
"""
import docx
import openpyxl
import pytest
from pptx import Presentation

from app.ai_chat.services.document_loader import load_document_text


def test_extract_txt(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("Xin chao lop hoc.", encoding="utf-8")

    text = load_document_text(str(path))
    assert "Xin chao lop hoc." in text


def test_extract_docx_marks_headings_as_markdown(tmp_path):
    doc = docx.Document()
    doc.add_heading("Chuong 1", level=1)
    doc.add_paragraph("Noi dung chuong 1.")
    path = tmp_path / "sample.docx"
    doc.save(str(path))

    text = load_document_text(str(path))
    assert "# Chuong 1" in text
    assert "Noi dung chuong 1." in text


def test_extract_pptx_marks_slide_number(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide dau tien"
    path = tmp_path / "sample.pptx"
    prs.save(str(path))

    text = load_document_text(str(path))
    assert "## Slide 1" in text
    assert "Slide dau tien" in text


def test_extract_xlsx_marks_sheet_name(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diem"
    ws.append(["Ten", "Diem"])
    ws.append(["An", 9.5])
    path = tmp_path / "sample.xlsx"
    wb.save(str(path))

    text = load_document_text(str(path))
    assert "## Sheet: Diem" in text
    assert "An" in text and "9.5" in text


def test_unsupported_extension_raises_value_error(tmp_path):
    path = tmp_path / "sample.zip"
    path.write_bytes(b"PK\x03\x04fake-zip-bytes")

    with pytest.raises(ValueError):
        load_document_text(str(path))


def test_tcvn3_mis_encoded_text_gets_fixed_to_unicode(tmp_path):
    """Mo phong van ban PDF/DOC cu bi doc sai bang ma TCVN3 (vd: 'chµo' -> 'chào')."""
    mis_encoded = "Xin chµo c¸c em häc sinh, chµo mõng b¹n ®Õn víi líp häc"
    path = tmp_path / "old.txt"
    path.write_text(mis_encoded, encoding="utf-8")

    text = load_document_text(str(path))
    assert "chào" in text
